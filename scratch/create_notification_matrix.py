import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Output path
target_dir = "/Users/congnguyen/Library/CloudStorage/OneDrive-Personal/Documents/Synodus/GOV/01.VNA/VNA/images/document/BA DOC"
os.makedirs(target_dir, exist_ok=True)
output_path = os.path.join(target_dir, "Danh_Sach_Kich_Ban_Thong_Bao_Notification_VNA.xlsx")

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Color Palette (Vietnam Airlines Professional Style)
NAVY_HEADER = "005F6E"     # VNA Deep Teal
GOLD_ACCENT = "D4A017"     # VNA Lotus Gold
LIGHT_BLUE_BG = "E6F4F6"   # Soft Teal
ZEBRA_BG = "F8FAFC"        # Light gray-blue
WHITE = "FFFFFF"
BORDER_GRAY = "D1D5DB"
HEADER_BORDER = "003B46"

# Fonts
font_title = Font(name="Segoe UI", size=16, bold=True, color="003B46")
font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
font_section = Font(name="Segoe UI", size=12, bold=True, color="005F6E")
font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
font_data = Font(name="Segoe UI", size=9, color="1F2937")
font_data_bold = Font(name="Segoe UI", size=9, bold=True, color="1F2937")
font_badge = Font(name="Segoe UI", size=8.5, bold=True)

# Borders
thin_side = Side(border_style="thin", color=BORDER_GRAY)
cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
header_border = Border(left=thin_side, right=thin_side, top=Side(border_style="medium", color=HEADER_BORDER), bottom=Side(border_style="medium", color=HEADER_BORDER))

# Alignments
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

print("Starting generation script...")
