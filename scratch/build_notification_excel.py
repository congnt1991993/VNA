# -*- coding: utf-8 -*-
import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

output_dir = "/Users/congnguyen/Library/CloudStorage/OneDrive-Personal/Documents/Synodus/GOV/01.VNA/VNA/images/document/BA DOC"
os.makedirs(output_dir, exist_ok=True)
excel_file = os.path.join(output_dir, "Danh_Sach_Kich_Ban_Thong_Bao_Notification_VNA.xlsx")

wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

# Color Palette
COLOR_NAVY_DARK = "003B46"     # Deep Teal / Navy
COLOR_NAVY_LIGHT = "005F6E"    # VNA Teal Header
COLOR_GOLD = "D4A017"          # Lotus Gold
COLOR_ZEBRA = "F8FAFC"         # Light Gray-Blue
COLOR_WHITE = "FFFFFF"
COLOR_GRAY_BORDER = "D1D5DB"

# Fills & Fonts for Badges
FILL_REQ = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Red
FONT_REQ = Font(name="Segoe UI", size=8.5, bold=True, color="991B1B")

FILL_WARN = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Amber
FONT_WARN = Font(name="Segoe UI", size=8.5, bold=True, color="92400E")

FILL_SUCC = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Green
FONT_SUCC = Font(name="Segoe UI", size=8.5, bold=True, color="166534")

FILL_INFO = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid") # Blue
FONT_INFO = Font(name="Segoe UI", size=8.5, bold=True, color="1E40AF")

FILL_ORIGIN_CORE = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid") # Sky blue
FONT_ORIGIN_CORE = Font(name="Segoe UI", size=8.5, bold=True, color="0369A1")

FILL_ORIGIN_PROP = PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid") # Purple light
FONT_ORIGIN_PROP = Font(name="Segoe UI", size=8.5, bold=True, color="6B21A8")

# Common Fonts
font_title = Font(name="Segoe UI", size=13.5, bold=True, color="003B46")
font_subtitle = Font(name="Segoe UI", size=9.5, italic=True, color="4B5563")
font_tbl_header = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
font_data_cell = Font(name="Segoe UI", size=9, color="1F2937")
font_code_cell = Font(name="Consolas", size=9, bold=True, color="005F6E")
font_bold_cell = Font(name="Segoe UI", size=9, bold=True, color="111827")

# Borders
border_thin = Side(border_style="thin", color=COLOR_GRAY_BORDER)
border_cell = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

# Fills
fill_header = PatternFill(start_color=COLOR_NAVY_LIGHT, end_color=COLOR_NAVY_LIGHT, fill_type="solid")
fill_zebra = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid")
fill_white = PatternFill(start_color=COLOR_WHITE, end_color=COLOR_WHITE, fill_type="solid")

def apply_title_block(ws, title_text, subtitle_text, col_span=14):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_span)
    ws.cell(1, 1).value = "TỔNG CÔNG TY HÀNG KHÔNG VIỆT NAM (VIETNAM AIRLINES) - CỔNG THÔNG TIN NET ZERO & BÁO CÁO PHÁT TRIỂN BỀN VỮNG"
    ws.cell(1, 1).font = Font(name="Segoe UI", size=8.5, bold=True, color="6B7280")
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_span)
    ws.cell(2, 1).value = title_text
    ws.cell(2, 1).font = font_title
    ws.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=col_span)
    ws.cell(3, 1).value = subtitle_text
    ws.cell(3, 1).font = font_subtitle
    ws.cell(3, 1).alignment = Alignment(horizontal="left", vertical="center")
    
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 8

# Master Scenarios Dataset
ALL_SCENARIOS = [
    # 1. CORE SCENARIO 1
    {
        "id": "NOTIF-CMS-001",
        "module": "CMS - Quản lý Biểu đồ Công bố",
        "origin": "Yêu Cầu Trọng Tâm (Gốc)",
        "scenario": "Tổ ban gửi yêu cầu công bố phiên bản số liệu từ chức năng Điều chỉnh số liệu",
        "trigger": "Chuyên viên Tổ ban nhấn nút \"Đề nghị công bố\" trên dòng phiên bản tại trang Điều chỉnh số liệu công bố",
        "actor": "Chuyên viên Tổ ban (Tổ Khai thác, Ban ATCL, Tổ Kỹ thuật, Ban TCNL...)",
        "recipient": "Ban Quản trị CMS (CMS Admin), Lãnh đạo Ban KHPT phụ trách duyệt",
        "title": "[Yêu cầu công bố số liệu] {indicator_code} - {indicator_name}",
        "template": "{user_name} ({department}) vừa gửi đề nghị công bố phiên bản {version_number} cho biểu đồ/chỉ tiêu \"{chart_name}\". Vui lòng kiểm tra và duyệt xuất bản trên CMS.",
        "type": "Yêu cầu xử lý (Action Required)",
        "deeplink": "Điều hướng đến CMS -> Quản lý Biểu đồ Công bố -> Tab Backlog (Highlight thẻ biểu đồ)",
        "channels": "In-app Quả chuông, Toast popup, Email duyệt",
        "priority": "Cao (P1)"
    },
    # 2. CORE SCENARIO 2
    {
        "id": "NOTIF-CMS-002",
        "module": "CMS - Kho tài liệu PTBV",
        "origin": "Yêu Cầu Trọng Tâm (Gốc)",
        "scenario": "Tổ ban thêm tài liệu mới vào Kho tài liệu PTBV",
        "trigger": "Chuyên viên tải tệp lên và nhấn nút \"Thêm tài liệu\" / \"Lưu tài liệu\" tại Kho tài liệu PTBV chung",
        "actor": "Chuyên viên Tổ ban nghiệp vụ (Uploader)",
        "recipient": "Ban Quản trị CMS, Ban KHPT, Cán bộ các Tổ ban trong phạm vi tài liệu",
        "title": "[Kho tài liệu PTBV] Tài liệu mới được tải lên: {doc_name}",
        "template": "{user_name} ({department}) vừa thêm tài liệu mới \"{doc_name}\" (Mã: {doc_code}, Loại: {doc_type}, Phạm vi: {scope}) vào Kho tài liệu PTBV chung.",
        "type": "Thông tin (Info)",
        "deeplink": "Điều hướng đến Kho tài liệu PTBV chung -> Mở modal xem chi tiết tài liệu {doc_code}",
        "channels": "In-app Quả chuông, Toast popup",
        "priority": "Cao (P1)"
    },
    # 3. CORE SCENARIO 3
    {
        "id": "NOTIF-CMS-003",
        "module": "CMS - Kho tài liệu PTBV",
        "origin": "Yêu Cầu Trọng Tâm (Gốc)",
        "scenario": "Tổ ban xóa tài liệu khỏi Kho tài liệu PTBV",
        "trigger": "Người dùng có quyền nhấn biểu tượng Thùng rác và xác nhận xóa tài liệu khỏi danh sách",
        "actor": "Người sở hữu tài liệu / Quản trị viên Kho tài liệu",
        "recipient": "Ban Quản trị CMS, Ban KHPT, và người đã đăng tài liệu gốc (nếu Admin xóa)",
        "title": "[Kho tài liệu PTBV] Tài liệu đã bị xóa: {doc_name}",
        "template": "Tài liệu \"{doc_name}\" (Mã: {doc_code}) đã được gỡ bỏ khỏi Kho tài liệu PTBV chung bởi {user_name} ({department}).",
        "type": "Cảnh báo (Warning)",
        "deeplink": "Điều hướng đến Kho tài liệu PTBV chung -> Danh sách tài liệu hiện hành",
        "channels": "In-app Quả chuông, Audit Log",
        "priority": "Cao (P1)"
    },
    # 4. CORE SCENARIO 4
    {
        "id": "NOTIF-ADJ-001",
        "module": "Điều chỉnh số liệu công bố",
        "origin": "Yêu Cầu Trọng Tâm (Gốc)",
        "scenario": "Kéo thẻ yêu cầu từ Backlog lên Biểu đồ xuất bản (Sprint) tại CMS",
        "trigger": "Quản trị viên CMS kéo thả thẻ biểu đồ từ khu vực Backlog vào danh sách Biểu đồ xuất bản (Sprint)",
        "actor": "Quản trị viên CMS (CMS Admin / Ban KHPT)",
        "recipient": "Chuyên viên và Lãnh đạo Tổ ban phụ trách chỉ tiêu/biểu đồ được xuất bản",
        "title": "[Xuất bản Biểu đồ] Biểu đồ {chart_name} đã được chọn xuất bản",
        "template": "Biểu đồ \"{chart_name}\" (Chỉ tiêu {indicator_code}) của {department} đã được {user_name} đưa từ Backlog lên danh sách Biểu đồ xuất bản chính thức trên cổng thông tin Net Zero.",
        "type": "Thành công (Success)",
        "deeplink": "Điều hướng đến Điều chỉnh số liệu công bố -> Chọn chỉ tiêu {indicator_code} -> Badge \"✓ Đang công bố\"",
        "channels": "In-app Quả chuông, Toast popup",
        "priority": "Cao (P1)"
    },
    # 5. CORE SCENARIO 5
    {
        "id": "NOTIF-ADJ-002",
        "module": "Điều chỉnh số liệu công bố",
        "origin": "Yêu Cầu Trọng Tâm (Gốc)",
        "scenario": "Xóa thẻ yêu cầu trên Biểu đồ xuất bản tại CMS (Gỡ khỏi xuất bản)",
        "trigger": "Quản trị viên CMS nhấn biểu tượng Xóa (Thùng rác) tại dòng biểu đồ trên bảng Biểu đồ xuất bản",
        "actor": "Quản trị viên CMS (CMS Admin / Ban KHPT)",
        "recipient": "Tổ ban phụ trách chỉ tiêu/biểu đồ, Người tạo phiên bản công bố",
        "title": "[Gỡ xuất bản Biểu đồ] Biểu đồ {chart_name} đã được gỡ khỏi Trang chi tiết",
        "template": "Biểu đồ \"{chart_name}\" (Chỉ tiêu {indicator_code}) đã được {user_name} gỡ khỏi danh sách Biểu đồ xuất bản và chuyển về Backlog. Trạng thái công bố chuyển sang Tạm ẩn.",
        "type": "Cảnh báo (Warning)",
        "deeplink": "Điều hướng đến Điều chỉnh số liệu công bố -> Chọn chỉ tiêu {indicator_code} -> Tab Quản lý phiên bản",
        "channels": "In-app Quả chuông, Toast popup",
        "priority": "Cao (P1)"
    },
    # 6. PROPOSED 1: Kích hoạt phiên bản Active
    {
        "id": "NOTIF-ADJ-003",
        "module": "Điều chỉnh số liệu công bố",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Phê duyệt kích hoạt phiên bản chính thức (Active Version)",
        "trigger": "Quản trị viên hoặc Lãnh đạo phê duyệt kích hoạt một phiên bản số liệu sang trạng thái Active (thay thế bản cũ)",
        "actor": "Admin / Lãnh đạo Ban KHPT",
        "recipient": "Toàn bộ thành viên Tổ ban phụ trách chỉ tiêu, Ban Biên tập báo cáo thường niên",
        "title": "[Kích hoạt phiên bản] Đã kích hoạt bản công bố {version_number} - {chart_name}",
        "template": "Phiên bản \"{version_number} - {version_name}\" của biểu đồ \"{chart_name}\" đã được {user_name} kích hoạt làm bản số liệu công bố chính thức đối ngoại.",
        "type": "Thành công (Success)",
        "deeplink": "Điều hướng đến Điều chỉnh số liệu công bố -> Focus vào phiên bản Active của biểu đồ {chart_code}",
        "channels": "In-app Quả chuông, Email thông báo",
        "priority": "Cao (P1)"
    },
    # 7. PROPOSED 2: Yêu cầu gỡ công bố phiên bản
    {
        "id": "NOTIF-ADJ-004",
        "module": "Điều chỉnh số liệu công bố",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Tổ ban gửi yêu cầu khẩn đề nghị gỡ công bố phiên bản (Unpublish Request)",
        "trigger": "Tổ ban phát hiện số liệu cần đính chính khẩn cấp và nhấn \"Đề nghị gỡ công bố\" phiên bản đang Active",
        "actor": "Chuyên viên Tổ ban phụ trách",
        "recipient": "Quản trị viên CMS, Lãnh đạo Ban KHPT",
        "title": "[Yêu cầu gỡ công bố] Đề nghị gỡ phiên bản {version_number} - {chart_name}",
        "template": "{user_name} ({department}) đã gửi yêu cầu khẩn đề nghị gỡ công bố phiên bản {version_number} của biểu đồ \"{chart_name}\". Lý do: \"{reason}\".",
        "type": "Yêu cầu xử lý (Action Required)",
        "deeplink": "Điều hướng đến Điều chỉnh số liệu công bố -> Mở danh sách phiên bản của biểu đồ {chart_code}",
        "channels": "In-app Quả chuông, Toast khẩn cấp, Email",
        "priority": "Cao (P1)"
    },
    # 8. PROPOSED 3: Từ chối yêu cầu công bố
    {
        "id": "NOTIF-ADJ-005",
        "module": "Điều chỉnh số liệu công bố",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Từ chối yêu cầu công bố phiên bản số liệu kèm lý do",
        "trigger": "Quản trị viên CMS hoặc Lãnh đạo từ chối đề nghị công bố phiên bản và nhập lý do phản hồi",
        "actor": "Quản trị viên CMS / Lãnh đạo Ban KHPT",
        "recipient": "Người gửi đề nghị công bố (Chuyên viên tổ ban)",
        "title": "[Từ chối công bố] Phiên bản {version_number} - {chart_name} chưa được duyệt",
        "template": "Yêu cầu công bố phiên bản {version_number} cho biểu đồ \"{chart_name}\" đã bị từ chối bởi {user_name}. Lý do: \"{reject_reason}\". Vui lòng kiểm tra và hiệu chỉnh lại.",
        "type": "Cảnh báo (Warning)",
        "deeplink": "Điều hướng đến Điều chỉnh số liệu công bố -> Mở modal Điều chỉnh số liệu của phiên bản {version_number}",
        "channels": "In-app Quả chuông, Toast popup",
        "priority": "Cao (P1)"
    },
    # 9. PROPOSED 4: Cảnh báo sai lệch số liệu nguồn (Data Drift)
    {
        "id": "NOTIF-ADJ-006",
        "module": "Điều chỉnh số liệu công bố",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Cảnh báo sai lệch số liệu nguồn so với bản đã công bố (Data Drift Alert)",
        "trigger": "Hệ thống tự động quét phát hiện số liệu tại Form nhập liệu gốc có sự thay đổi sai lệch so với bản công bố Active",
        "actor": "Hệ thống tự động (Automated Audit Job)",
        "recipient": "Chuyên viên phụ trách chỉ tiêu, Quản trị viên số liệu",
        "title": "[Cảnh báo sai lệch số liệu] Chỉ tiêu {indicator_code} phát sinh thay đổi nguồn",
        "template": "Phát hiện số liệu gốc kỳ {period} của chỉ tiêu \"{indicator_name}\" vừa được cập nhật tại Form nhập liệu, có chênh lệch so với bản công bố {version_number}. Vui lòng rà soát lại.",
        "type": "Cảnh báo (Warning)",
        "deeplink": "Điều hướng đến Điều chỉnh số liệu công bố -> Bảng đối chiếu số liệu hệ thống vs số liệu điều chỉnh",
        "channels": "In-app Quả chuông, Email cảnh báo",
        "priority": "Trung bình (P2)"
    },
    # 10. PROPOSED 5: Cập nhật thuyết minh công bố
    {
        "id": "NOTIF-ADJ-007",
        "module": "Điều chỉnh số liệu công bố",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Cập nhật nội dung Thuyết minh công bố định tính của biểu đồ",
        "trigger": "Người dùng chỉnh sửa đoạn văn bản thuyết minh định tính (GRI Disclosure text) và lưu phiên bản mới",
        "actor": "Chuyên viên Ban Truyền thông / Ban KHPT",
        "recipient": "Các tổ ban liên quan, Ban biên tập Báo cáo ESG",
        "title": "[Cập nhật Thuyết minh] Biểu đồ {chart_name} có thuyết minh mới",
        "template": "{user_name} vừa cập nhật nội dung thuyết minh công bố cho biểu đồ \"{chart_name}\" (Chỉ tiêu {indicator_code}) theo phiên bản {version_number}.",
        "type": "Thông tin (Info)",
        "deeplink": "Điều hướng đến Điều chỉnh số liệu công bố -> Mở khung thuyết minh của chỉ tiêu {indicator_code}",
        "channels": "In-app Quả chuông",
        "priority": "Trung bình (P2)"
    },
    # 11. PROPOSED 6: Nhắc nhở định kỳ
    {
        "id": "NOTIF-ADJ-008",
        "module": "Điều chỉnh số liệu công bố",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Nhắc nhở định kỳ rà soát & cập nhật phiên bản số liệu công bố",
        "trigger": "Hệ thống tự động kích hoạt vào ngày 25 hàng tháng hoặc cuối quý theo lịch chốt kỳ báo cáo ESG",
        "actor": "Hệ thống tự động (Scheduled Reminder)",
        "recipient": "Chuyên viên và Đầu mối phụ trách chỉ tiêu của tất cả các Tổ ban",
        "title": "[Nhắc nhở định kỳ] Rà soát và cập nhật số liệu công bố kỳ {period}",
        "template": "Đã đến kỳ rà soát dữ liệu công bố {period}. Vui lòng kiểm tra số liệu các chỉ tiêu được phân công và tạo phiên bản công bố mới trước ngày {deadline}.",
        "type": "Yêu cầu xử lý (Action Required)",
        "deeplink": "Điều hướng đến Điều chỉnh số liệu công bố -> Tự động lọc chỉ tiêu của phòng ban đang đăng nhập",
        "channels": "In-app Quả chuông, Email nhắc nhở",
        "priority": "Cao (P1)"
    },
    # 12. PROPOSED 7: Sắp xếp lại thứ tự biểu đồ
    {
        "id": "NOTIF-CMS-004",
        "module": "CMS - Quản lý Biểu đồ Công bố",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Sắp xếp lại thứ tự hiển thị biểu đồ trên Trang chi tiết Trụ cột",
        "trigger": "Quản trị viên CMS kéo thả thay đổi vị trí thứ tự STT (#1, #2, #3...) của các biểu đồ trong Sprint",
        "actor": "Quản trị viên CMS",
        "recipient": "Nhóm quản trị CMS, Trưởng nhóm biên tập nội dung",
        "title": "[Cấu hình CMS] Đã cập nhật thứ tự hiển thị biểu đồ Trụ cột {pillar_name}",
        "template": "{user_name} vừa điều chỉnh lại thứ tự hiển thị của {count} biểu đồ thuộc Trụ cột {pillar_name} trên trang chi tiết công bố.",
        "type": "Thông tin (Info)",
        "deeplink": "Điều hướng đến CMS -> Quản lý Biểu đồ Công bố -> Xem danh sách biểu đồ Sprint",
        "channels": "In-app Quả chuông",
        "priority": "Thấp (P3)"
    },
    # 13. PROPOSED 8: Xuất bản Portal Net Zero
    {
        "id": "NOTIF-CMS-005",
        "module": "CMS - Xuất bản Portal Net Zero",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Kích hoạt xuất bản toàn diện Cổng thông tin Net Zero / Báo cáo ESG",
        "trigger": "Lãnh đạo Tổng công ty / CMS Admin nhấn \"Xuất bản phiên bản Portal\" đẩy toàn bộ cấu hình ra public",
        "actor": "Tổng Biên tập / Lãnh đạo phê duyệt",
        "recipient": "Toàn thể cán bộ nhân viên và người dùng hệ thống ESG VNA (All Users)",
        "title": "[Thông báo xuất bản] Cổng thông tin Net Zero đã công bố phiên bản mới {portal_version}",
        "template": "Cổng thông tin Net Zero & Báo cáo Phát triển bền vững Vietnam Airlines đã chính thức xuất bản phiên bản {portal_version}. Nhấn để xem trang chủ công bố.",
        "type": "Thành công (Success)",
        "deeplink": "Mở Landing Page công bố đối ngoại (Trang chủ Net Zero Portal)",
        "channels": "In-app Quả chuông, Banner trang chủ, Email",
        "priority": "Cao (P1)"
    },
    # 14. PROPOSED 9: Cảnh báo biểu đồ thiếu dữ liệu
    {
        "id": "NOTIF-CMS-006",
        "module": "CMS - Quản lý Biểu đồ Công bố",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Cảnh báo Biểu đồ xuất bản bị khuyết thiếu dữ liệu kỳ báo cáo mới",
        "trigger": "Hệ thống tự động phát hiện biểu đồ trong danh sách xuất bản nhưng chưa có số liệu kỳ mới nhất hoặc bản Active bị thiếu",
        "actor": "Hệ thống tự động (CMS Validator)",
        "recipient": "CMS Admin, Tổ ban phụ trách biểu đồ",
        "title": "[Cảnh báo CMS] Biểu đồ {chart_name} chưa có dữ liệu công bố kỳ {period}",
        "template": "Biểu đồ \"{chart_name}\" đang được đính kèm xuất bản nhưng chưa có dữ liệu kỳ {period} hoặc chưa có phiên bản Active. Vui lòng bổ sung trước khi xuất bản.",
        "type": "Cảnh báo (Warning)",
        "deeplink": "Điều hướng đến CMS -> Quản lý Biểu đồ -> Mở khung Preview biểu đồ bị thiếu",
        "channels": "In-app Quả chuông, Cảnh báo đỏ tại bảng CMS",
        "priority": "Cao (P1)"
    },
    # 15. PROPOSED 10: Phê duyệt tài liệu thành công
    {
        "id": "NOTIF-DOC-004",
        "module": "CMS - Kho tài liệu PTBV",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Phê duyệt tài liệu đưa vào Kho lưu trữ chung thành công",
        "trigger": "Người kiểm duyệt nhấn \"Phê duyệt\" tài liệu tại màn hình Phê duyệt tài liệu (DocumentApproval.tsx)",
        "actor": "Người kiểm duyệt tài liệu (Ban KHPT / Admin)",
        "recipient": "Người tải lên tài liệu (Uploader)",
        "title": "[Phê duyệt tài liệu] Tài liệu \"{doc_name}\" đã được duyệt vào Kho chung",
        "template": "Tài liệu \"{doc_name}\" (Mã: {doc_code}) gửi ngày {submit_date} đã được {approver_name} phê duyệt thành công và đưa vào Kho tài liệu PTBV chung.",
        "type": "Thành công (Success)",
        "deeplink": "Điều hướng đến Kho tài liệu PTBV -> Xem tài liệu đã duyệt trong kho chung",
        "channels": "In-app Quả chuông, Email xác nhận",
        "priority": "Cao (P1)"
    },
    # 16. PROPOSED 11: Từ chối duyệt tài liệu
    {
        "id": "NOTIF-DOC-005",
        "module": "CMS - Kho tài liệu PTBV",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Từ chối phê duyệt tài liệu đưa vào Kho lưu trữ chung",
        "trigger": "Người kiểm duyệt nhấn \"Từ chối\" kèm lý do phản hồi tại màn hình Phê duyệt tài liệu",
        "actor": "Người kiểm duyệt tài liệu (Ban KHPT / Admin)",
        "recipient": "Người tải lên tài liệu (Uploader)",
        "title": "[Từ chối tài liệu] Tài liệu \"{doc_name}\" chưa được duyệt vào Kho chung",
        "template": "Yêu cầu lưu kho tài liệu \"{doc_name}\" (Mã: {doc_code}) đã bị từ chối bởi {approver_name}. Lý do: \"{reject_reason}\". Vui lòng kiểm tra và gửi lại.",
        "type": "Cảnh báo (Warning)",
        "deeplink": "Điều hướng đến Kho tài liệu PTBV -> Tab Yêu cầu tải lên (Xem lý do từ chối)",
        "channels": "In-app Quả chuông, Toast cảnh báo",
        "priority": "Cao (P1)"
    },
    # 17. PROPOSED 12: Cập nhật phiên bản mới tài liệu
    {
        "id": "NOTIF-DOC-006",
        "module": "CMS - Kho tài liệu PTBV",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Cập nhật phiên bản mới của tài liệu trong Kho PTBV (v1.0 -> v2.0)",
        "trigger": "Người dùng tải tệp thay thế hoặc cập nhật phiên bản tài liệu mới (Semantic Versioning)",
        "actor": "Người sở hữu tài liệu / Quản trị viên Kho",
        "recipient": "Toàn bộ cán bộ nhân viên có quyền truy cập tài liệu này",
        "title": "[Cập nhật phiên bản] Tài liệu \"{doc_name}\" đã nâng cấp lên bản {doc_version}",
        "template": "Tài liệu \"{doc_name}\" vừa được {user_name} cập nhật phiên bản mới {doc_version} (Ghi chú: {version_notes}). Nhấn để tải về tài liệu mới nhất.",
        "type": "Thông tin (Info)",
        "deeplink": "Điều hướng đến Kho tài liệu PTBV -> Mở modal Chi tiết tài liệu với bản mới nhất",
        "channels": "In-app Quả chuông",
        "priority": "Trung bình (P2)"
    },
    # 18. PROPOSED 13: Xin quyền truy cập tài liệu nội bộ
    {
        "id": "NOTIF-DOC-007",
        "module": "CMS - Kho tài liệu PTBV",
        "origin": "Phân Tích Đề Xuất Bổ Sung",
        "scenario": "Yêu cầu cấp quyền truy cập tài liệu nội bộ / bảo mật",
        "trigger": "Người dùng nhấn \"Gửi yêu cầu truy cập\" vào tài liệu được đánh dấu Nội bộ",
        "actor": "Người dùng có nhu cầu truy cập",
        "recipient": "Người sở hữu tài liệu (Uploader) và Quản trị viên hệ thống",
        "title": "[Yêu cầu truy cập] Cán bộ {user_name} xin quyền xem tài liệu {doc_name}",
        "template": "{user_name} ({department}) vừa gửi yêu cầu xin quyền xem/tải tài liệu nội bộ \"{doc_name}\" (Mã: {doc_code}). Vui lòng phê duyệt cấp quyền.",
        "type": "Yêu cầu xử lý (Action Required)",
        "deeplink": "Điều hướng đến Kho tài liệu PTBV -> Quản lý phân quyền tài liệu",
        "channels": "In-app Quả chuông, Toast popup, Email",
        "priority": "Trung bình (P2)"
    }
]

# =========================================================================
# SHEET 1: DANH MỤC TỔNG HỢP TOÀN BỘ KỊCH BẢN
# =========================================================================
ws1 = wb.create_sheet(title="1. Danh Mục Tổng Hợp")
apply_title_block(
    ws1,
    title_text="MA TRẬN KỊCH BẢN THÔNG BÁO QUẢ CHUÔNG (NOTIFICATION CATALOG)",
    subtitle_text="Đặc tả kịch bản thông báo In-app Quả chuông cho phân hệ CMS và phân hệ Điều chỉnh số liệu công bố - Vietnam Airlines",
    col_span=14
)

headers1 = [
    "STT", "Mã Kịch Bản", "Phân Hệ Chức Năng", "Nguồn Gốc Kịch Bản",
    "Tình Huống Nghiệp Vụ", "Hành Động Kích Hoạt (Trigger Event)", "Người Thực Hiện (Actor)",
    "Đối Tượng Nhận Thông Báo", "Tiêu Đề Thông Báo Quả Chuông", "Nội Dung Thông Báo Chi Tiết (Message Template)",
    "Phân Loại Mức Độ", "Điều Hướng Khi Nhấp (Deep Link)", "Kênh Bổ Trợ", "Ưu Tiên"
]

ws1.row_dimensions[5].height = 28
for col_idx, h_text in enumerate(headers1, 1):
    cell = ws1.cell(5, col_idx)
    cell.value = h_text
    cell.font = font_tbl_header
    cell.fill = fill_header
    cell.border = border_cell
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for r_idx, sc in enumerate(ALL_SCENARIOS, 6):
    is_even = (r_idx % 2 == 0)
    row_fill = fill_zebra if is_even else fill_white
    ws1.row_dimensions[r_idx].height = 42
    
    # 1. STT
    c1 = ws1.cell(r_idx, 1, r_idx - 5)
    c1.alignment = Alignment(horizontal="center", vertical="center")
    
    # 2. Mã Kịch Bản
    c2 = ws1.cell(r_idx, 2, sc["id"])
    c2.font = font_code_cell
    c2.alignment = Alignment(horizontal="center", vertical="center")
    
    # 3. Phân Hệ
    c3 = ws1.cell(r_idx, 3, sc["module"])
    c3.font = font_bold_cell
    c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # 4. Nguồn Gốc
    c4 = ws1.cell(r_idx, 4, sc["origin"])
    if "Gốc" in sc["origin"]:
        c4.fill = FILL_ORIGIN_CORE
        c4.font = FONT_ORIGIN_CORE
    else:
        c4.fill = FILL_ORIGIN_PROP
        c4.font = FONT_ORIGIN_PROP
    c4.alignment = Alignment(horizontal="center", vertical="center")
    
    # 5. Tình Huống
    c5 = ws1.cell(r_idx, 5, sc["scenario"])
    c5.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # 6. Thao tác Trigger
    c6 = ws1.cell(r_idx, 6, sc["trigger"])
    c6.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # 7. Người thực hiện
    c7 = ws1.cell(r_idx, 7, sc["actor"])
    c7.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # 8. Người nhận
    c8 = ws1.cell(r_idx, 8, sc["recipient"])
    c8.font = font_bold_cell
    c8.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # 9. Tiêu đề
    c9 = ws1.cell(r_idx, 9, sc["title"])
    c9.font = font_bold_cell
    c9.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # 10. Nội dung template
    c10 = ws1.cell(r_idx, 10, sc["template"])
    c10.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # 11. Mức độ
    c11 = ws1.cell(r_idx, 11, sc["type"])
    c11.alignment = Alignment(horizontal="center", vertical="center")
    if "Action Required" in sc["type"]:
        c11.fill = FILL_REQ
        c11.font = FONT_REQ
    elif "Warning" in sc["type"]:
        c11.fill = FILL_WARN
        c11.font = FONT_WARN
    elif "Success" in sc["type"]:
        c11.fill = FILL_SUCC
        c11.font = FONT_SUCC
    else:
        c11.fill = FILL_INFO
        c11.font = FONT_INFO
        
    # 12. Deep Link
    c12 = ws1.cell(r_idx, 12, sc["deeplink"])
    c12.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # 13. Kênh
    c13 = ws1.cell(r_idx, 13, sc["channels"])
    c13.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # 14. Ưu tiên
    c14 = ws1.cell(r_idx, 14, sc["priority"])
    c14.alignment = Alignment(horizontal="center", vertical="center")
    if "P1" in sc["priority"]:
        c14.font = Font(name="Segoe UI", size=9, bold=True, color="991B1B")
    elif "P2" in sc["priority"]:
        c14.font = Font(name="Segoe UI", size=9, bold=True, color="D97706")
    else:
        c14.font = Font(name="Segoe UI", size=9, color="4B5563")

    for col in range(1, 15):
        c = ws1.cell(r_idx, col)
        c.border = border_cell
        if col not in [4, 11]: # Don't overwrite colored badges
            if c.font == font_data_cell or c.font is None:
                c.font = font_data_cell
            c.fill = row_fill

# Enable auto filter
ws1.auto_filter.ref = f"A5:N{len(ALL_SCENARIOS) + 5}"
ws1.freeze_panes = "C6"

# =========================================================================
# SHEET 2: KỊCH BẢN YÊU CẦU GỐC (TRỌNG TÂM NGƯỜI DÙNG NÊU)
# =========================================================================
ws2 = wb.create_sheet(title="2. Kịch Bản Yêu Cầu Gốc")
apply_title_block(
    ws2,
    title_text="KỊCH BẢN THÔNG BÁO THEO YÊU CẦU TRỌNG TÂM CỦA NGƯỜI DÙNG",
    subtitle_text="Chi tiết 5 kịch bản thông báo chính thức tại phân hệ CMS & Điều chỉnh số liệu công bố theo yêu cầu của Vietnam Airlines",
    col_span=13
)

headers2 = [
    "STT", "Mã Kịch Bản", "Phân Hệ / Chức Năng", "Yêu Cầu Từ Người Dùng",
    "Sự Kiện Kích Hoạt (Trigger Event)", "Người Thực Hiện Thao Tác", "Đối Tượng Nhận Thông Báo",
    "Tiêu Đề Thông Báo Quả Chuông", "Nội Dung Hiển Thị Mẫu (In-app Notification)", "Mức Độ Thông Báo",
    "Hành Vi Khi Người Dùng Bấm Vào Thông Báo (Deep Link)", "Kênh Triển Khai", "Mức Độ Ưu Tiên"
]

ws2.row_dimensions[5].height = 28
for col_idx, h_text in enumerate(headers2, 1):
    cell = ws2.cell(5, col_idx)
    cell.value = h_text
    cell.font = font_tbl_header
    cell.fill = fill_header
    cell.border = border_cell
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

core_scenarios = [s for s in ALL_SCENARIOS if "Gốc" in s["origin"]]

for r_idx, sc in enumerate(core_scenarios, 6):
    is_even = (r_idx % 2 == 0)
    row_fill = fill_zebra if is_even else fill_white
    ws2.row_dimensions[r_idx].height = 46
    
    ws2.cell(r_idx, 1, r_idx - 5).alignment = Alignment(horizontal="center", vertical="center")
    c2 = ws2.cell(r_idx, 2, sc["id"])
    c2.font = font_code_cell
    c2.alignment = Alignment(horizontal="center", vertical="center")
    
    c3 = ws2.cell(r_idx, 3, sc["module"])
    c3.font = font_bold_cell
    c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    c4 = ws2.cell(r_idx, 4, sc["scenario"])
    c4.font = font_bold_cell
    c4.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    ws2.cell(r_idx, 5, sc["trigger"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws2.cell(r_idx, 6, sc["actor"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    c7 = ws2.cell(r_idx, 7, sc["recipient"])
    c7.font = font_bold_cell
    c7.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    c8 = ws2.cell(r_idx, 8, sc["title"])
    c8.font = font_bold_cell
    c8.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    ws2.cell(r_idx, 9, sc["template"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    c10 = ws2.cell(r_idx, 10, sc["type"])
    c10.alignment = Alignment(horizontal="center", vertical="center")
    if "Action Required" in sc["type"]:
        c10.fill = FILL_REQ; c10.font = FONT_REQ
    elif "Warning" in sc["type"]:
        c10.fill = FILL_WARN; c10.font = FONT_WARN
    elif "Success" in sc["type"]:
        c10.fill = FILL_SUCC; c10.font = FONT_SUCC
    else:
        c10.fill = FILL_INFO; c10.font = FONT_INFO
        
    ws2.cell(r_idx, 11, sc["deeplink"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws2.cell(r_idx, 12, sc["channels"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    c13 = ws2.cell(r_idx, 13, sc["priority"])
    c13.alignment = Alignment(horizontal="center", vertical="center")
    c13.font = Font(name="Segoe UI", size=9, bold=True, color="991B1B")

    for col in range(1, 14):
        c = ws2.cell(r_idx, col)
        c.border = border_cell
        if col != 10:
            if c.font == font_data_cell or c.font is None:
                c.font = font_data_cell
            c.fill = row_fill

ws2.auto_filter.ref = f"A5:M{len(core_scenarios) + 5}"
ws2.freeze_panes = "C6"

# =========================================================================
# SHEET 3: ĐỀ XUẤT BỔ SUNG & PHÂN TÍCH BA CHUYÊN SÂU
# =========================================================================
ws3 = wb.create_sheet(title="3. Đề Xuất Bổ Sung & Phân Tích")
apply_title_block(
    ws3,
    title_text="PHÂN TÍCH NGHIỆP VỤ & ĐỀ XUẤT BỔ SUNG KỊCH BẢN THÔNG BÁO",
    subtitle_text="Đánh giá rủi ro khoảng trống quy trình (Blind Spots), giá trị nghiệp vụ và giải pháp thông báo bổ trợ cho CMS & Điều chỉnh số liệu",
    col_span=10
)

headers3 = [
    "STT", "Mã Kịch Bản", "Phân Hệ Chức Năng", "Kịch Bản Nghiệp Vụ Bổ Sung",
    "Rủi Ro Nghiệp Vụ Nếu Thiếu Thông Báo Này (Risk / Pain Point)",
    "Giá Trị Mang Lại Cho Vietnam Airlines (Business Value)",
    "Hành Động Kích Hoạt & Đối Tượng Nhận Tin",
    "Nội Dung Thông Báo Đề Xuất",
    "Phân Loại Mức Độ", "Mức Độ Ưu Tiên Triển Khai"
]

ws3.row_dimensions[5].height = 28
for col_idx, h_text in enumerate(headers3, 1):
    cell = ws3.cell(5, col_idx)
    cell.value = h_text
    cell.font = font_tbl_header
    cell.fill = fill_header
    cell.border = border_cell
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

proposed_details = [
    {
        "id": "NOTIF-ADJ-003",
        "module": "Điều chỉnh số liệu công bố",
        "scenario": "Phê duyệt kích hoạt phiên bản chính thức (Active Version)",
        "risk": "Các tổ ban và bộ phận truyền thông không biết phiên bản nào đang là số liệu chính thức được phép trích xuất phục vụ kiểm toán đối ngoại.",
        "value": "Minh bạch hóa phiên bản dữ liệu hiện hành, đảm bảo tính duy nhất và nhất quán của số liệu công bố trên toàn hệ thống.",
        "trigger_recipient": "Admin kích hoạt bản Active -> Gửi toàn bộ Tổ ban phụ trách & Ban Biên tập báo cáo",
        "template": "Phiên bản \"{version_number} - {version_name}\" của biểu đồ \"{chart_name}\" đã được {user_name} kích hoạt làm bản số liệu công bố chính thức đối ngoại.",
        "type": "Thành công (Success)",
        "priority": "Cao (P1)"
    },
    {
        "id": "NOTIF-ADJ-004",
        "module": "Điều chỉnh số liệu công bố",
        "scenario": "Yêu cầu khẩn gỡ công bố phiên bản (Unpublish Request)",
        "risk": "Khi phát hiện số liệu sai sót nghiêm trọng hoặc có thay đổi kiểm toán, nếu không có cơ chế báo khẩn, số liệu sai sẽ tiếp tục hiển thị đối ngoại gây khủng hoảng thông tin.",
        "value": "Phản ứng nhanh với sự cố sai lệch số liệu, cho phép thu hồi và tạm ẩn biểu đồ ngay lập tức để rà soát.",
        "trigger_recipient": "Tổ ban gửi yêu cầu gỡ công bố -> Gửi Quản trị viên CMS & Lãnh đạo Ban KHPT",
        "template": "{user_name} ({department}) đã gửi yêu cầu khẩn đề nghị gỡ công bố phiên bản {version_number} của biểu đồ \"{chart_name}\". Lý do: \"{reason}\".",
        "type": "Yêu cầu xử lý (Action Required)",
        "priority": "Cao (P1)"
    },
    {
        "id": "NOTIF-ADJ-005",
        "module": "Điều chỉnh số liệu công bố",
        "scenario": "Từ chối yêu cầu công bố số liệu kèm lý do",
        "risk": "Chuyên viên tổ ban gửi yêu cầu công bố nhưng không nhận được phản hồi, không biết lý do bị chậm hoặc bị từ chối, gây tắc nghẽn tiến độ đóng kỳ báo cáo.",
        "value": "Khép kín chu trình kiểm duyệt (Approval Loop), giúp người nhập dữ liệu nhanh chóng sửa chữa đúng trọng tâm yêu cầu.",
        "trigger_recipient": "Admin/Lãnh đạo bấm Từ chối -> Gửi Chuyên viên tổ ban đã tạo yêu cầu",
        "template": "Yêu cầu công bố phiên bản {version_number} cho biểu đồ \"{chart_name}\" đã bị từ chối bởi {user_name}. Lý do: \"{reject_reason}\". Vui lòng kiểm tra và hiệu chỉnh lại.",
        "type": "Cảnh báo (Warning)",
        "priority": "Cao (P1)"
    },
    {
        "id": "NOTIF-ADJ-006",
        "module": "Điều chỉnh số liệu công bố",
        "scenario": "Cảnh báo sai lệch số liệu nguồn so với bản đã công bố (Data Drift Alert)",
        "risk": "Dữ liệu ở các phân hệ nguồn (Nhập liệu thủ công, ERP, Flight Data) bị ai đó sửa đổi ngầm sau khi đã công bố, dẫn đến số liệu công bố vênh với số liệu sổ sách kế toán.",
        "value": "Tự động phát hiện bất đồng bộ dữ liệu (Data Integrity Audit), bảo vệ tính pháp lý của báo cáo thường niên.",
        "trigger_recipient": "Hệ thống tự động phát hiện lệch -> Gửi Chuyên viên phụ trách chỉ tiêu & Quản trị viên số liệu",
        "template": "Phát hiện số liệu gốc kỳ {period} của chỉ tiêu \"{indicator_name}\" vừa được cập nhật tại Form nhập liệu, có chênh lệch so với bản công bố {version_number}. Vui lòng rà soát lại.",
        "type": "Cảnh báo (Warning)",
        "priority": "Trung bình (P2)"
    },
    {
        "id": "NOTIF-ADJ-007",
        "module": "Điều chỉnh số liệu công bố",
        "scenario": "Cập nhật nội dung Thuyết minh công bố định tính của biểu đồ",
        "risk": "Thuyết minh giải trình biến động (GRI Disclosure) bị cập nhật nhưng các bên liên quan không nắm được ngữ cảnh phân tích mới.",
        "value": "Đồng bộ hóa thông điệp truyền thông giữa số liệu biểu đồ và văn bản phân tích nguyên nhân tăng giảm phát thải.",
        "trigger_recipient": "Biên tập viên lưu thuyết minh -> Gửi các Tổ ban liên quan & Ban Truyền thông",
        "template": "{user_name} vừa cập nhật nội dung thuyết minh công bố cho biểu đồ \"{chart_name}\" (Chỉ tiêu {indicator_code}) theo phiên bản {version_number}.",
        "type": "Thông tin (Info)",
        "priority": "Trung bình (P2)"
    },
    {
        "id": "NOTIF-ADJ-008",
        "module": "Điều chỉnh số liệu công bố",
        "scenario": "Nhắc nhở định kỳ rà soát & cập nhật phiên bản số liệu công bố",
        "risk": "Các đơn vị chậm trễ nộp báo cáo định kỳ tháng/quý dẫn đến việc cổng Net Zero bị chậm tiến độ cập nhật số liệu mới cho cổ đông và nhà đầu tư.",
        "value": "Tự động hóa quy trình đôn đốc tiến độ, giảm thiểu thời gian nhân sự quản trị phải đi nhắc việc thủ công.",
        "trigger_recipient": "Hệ thống tự động nhắc ngày 25 hàng tháng -> Gửi Đầu mối nhập liệu của các Tổ ban",
        "template": "Đã đến kỳ rà soát dữ liệu công bố {period}. Vui lòng kiểm tra số liệu các chỉ tiêu được phân công và tạo phiên bản công bố mới trước ngày {deadline}.",
        "type": "Yêu cầu xử lý (Action Required)",
        "priority": "Cao (P1)"
    },
    {
        "id": "NOTIF-CMS-004",
        "module": "CMS - Quản lý Biểu đồ Công bố",
        "scenario": "Sắp xếp lại thứ tự hiển thị biểu đồ trên Trang chi tiết Trụ cột",
        "risk": "Thay đổi vị trí biểu đồ trọng điểm vô tình làm đảo lộn cấu trúc logic của báo cáo chuyên môn mà ban phụ trách trụ cột không hay biết.",
        "value": "Lưu vết và thông báo thay đổi layout giao diện, giúp ban phụ trách kiểm soát thứ tự ưu tiên truyền thông.",
        "trigger_recipient": "Admin kéo thả đổi thứ tự biểu đồ -> Gửi Ban phụ trách Trụ cột & Ban Truyền thông",
        "template": "{user_name} vừa điều chỉnh lại thứ tự hiển thị của {count} biểu đồ thuộc Trụ cột {pillar_name} trên trang chi tiết công bố.",
        "type": "Thông tin (Info)",
        "priority": "Thấp (P3)"
    },
    {
        "id": "NOTIF-CMS-005",
        "module": "CMS - Xuất bản Portal Net Zero",
        "scenario": "Kích hoạt xuất bản toàn diện Cổng thông tin Net Zero / Báo cáo ESG",
        "risk": "Hệ thống đã cập nhật bộ số liệu mới ra ngoài cổng thông tin công cộng nhưng nội bộ cán bộ công nhân viên không nắm được để truyền thông thống nhất.",
        "value": "Tạo sự đồng thuận cao, thông tin kịp thời đến toàn bộ Tổng công ty về thành tựu và chỉ số Net Zero mới nhất.",
        "trigger_recipient": "Lãnh đạo bấm Xuất bản Portal -> Gửi Toàn thể người dùng hệ thống ESG VNA",
        "template": "Cổng thông tin Net Zero & Báo cáo Phát triển bền vững Vietnam Airlines đã chính thức xuất bản phiên bản {portal_version}. Nhấn để xem trang chủ công bố.",
        "type": "Thành công (Success)",
        "priority": "Cao (P1)"
    },
    {
        "id": "NOTIF-CMS-006",
        "module": "CMS - Quản lý Biểu đồ Công bố",
        "scenario": "Cảnh báo Biểu đồ xuất bản bị khuyết thiếu dữ liệu kỳ báo cáo mới",
        "risk": "Biểu đồ được xuất bản công khai ra ngoài nhưng bị trắng số liệu (Blank chart) hoặc chỉ có dữ liệu cũ của năm ngoái, ảnh hưởng hình ảnh chuyên nghiệp của VNA.",
        "value": "Ngăn chặn lỗi hiển thị biểu đồ rỗng trên cổng thông tin đối ngoại, đóng vai trò như chốt kiểm soát chất lượng (Quality Gate).",
        "trigger_recipient": "Hệ thống phát hiện thiếu số liệu -> Gửi Quản trị viên CMS & Tổ ban liên quan",
        "template": "Biểu đồ \"{chart_name}\" đang được đính kèm xuất bản nhưng chưa có dữ liệu kỳ {period} hoặc chưa có phiên bản Active. Vui lòng bổ sung trước khi xuất bản.",
        "type": "Cảnh báo (Warning)",
        "priority": "Cao (P1)"
    },
    {
        "id": "NOTIF-DOC-004",
        "module": "CMS - Kho tài liệu PTBV",
        "scenario": "Phê duyệt tài liệu đưa vào Kho lưu trữ chung thành công",
        "risk": "Người đăng tải không biết tài liệu đã được duyệt hay chưa, dẫn đến việc tải lên trùng lặp hoặc không kịp thời chia sẻ cho đối tác.",
        "value": "Rút ngắn vòng quay tài liệu, giúp người nộp tài liệu an tâm rằng văn bản đã được pháp chế/Ban KHPT công nhận lưu kho.",
        "trigger_recipient": "Người kiểm duyệt bấm Phê duyệt -> Gửi Người đăng tải tài liệu (Uploader)",
        "template": "Tài liệu \"{doc_name}\" (Mã: {doc_code}) gửi ngày {submit_date} đã được {approver_name} phê duyệt thành công và đưa vào Kho tài liệu PTBV chung.",
        "type": "Thành công (Success)",
        "priority": "Cao (P1)"
    },
    {
        "id": "NOTIF-DOC-005",
        "module": "CMS - Kho tài liệu PTBV",
        "scenario": "Từ chối phê duyệt tài liệu đưa vào Kho lưu trữ chung",
        "risk": "Tài liệu bị treo không được duyệt mà không rõ nguyên nhân, dẫn đến thất lạc hồ sơ chứng nhận môi trường quan trọng.",
        "value": "Phản hồi minh bạch nguyên nhân từ chối (thiếu chữ ký số, sai định dạng, hết hạn) để đơn vị nhanh chóng hoàn thiện hồ sơ.",
        "trigger_recipient": "Người kiểm duyệt bấm Từ chối -> Gửi Người đăng tải tài liệu",
        "template": "Yêu cầu lưu kho tài liệu \"{doc_name}\" (Mã: {doc_code}) đã bị từ chối bởi {approver_name}. Lý do: \"{reject_reason}\". Vui lòng kiểm tra và gửi lại.",
        "type": "Cảnh báo (Warning)",
        "priority": "Cao (P1)"
    },
    {
        "id": "NOTIF-DOC-006",
        "module": "CMS - Kho tài liệu PTBV",
        "scenario": "Cập nhật phiên bản mới của tài liệu trong Kho PTBV (v1.0 -> v2.0)",
        "risk": "Cán bộ các ban tiếp tục sử dụng các quy chuẩn cũ, biểu mẫu cũ đã lỗi thời do không biết đã có bản cập nhật mới được ban hành.",
        "value": "Đảm bảo toàn bộ Tổng công ty luôn áp dụng phiên bản tài liệu chuẩn hóa mới nhất theo chuẩn quốc tế IATA/ICAO.",
        "trigger_recipient": "Người quản lý kho cập nhật phiên bản -> Gửi Toàn bộ cán bộ có quyền truy cập",
        "template": "Tài liệu \"{doc_name}\" vừa được {user_name} cập nhật phiên bản mới {doc_version} (Ghi chú: {version_notes}). Nhấn để tải về tài liệu mới nhất.",
        "type": "Thông tin (Info)",
        "priority": "Trung bình (P2)"
    },
    {
        "id": "NOTIF-DOC-007",
        "module": "CMS - Kho tài liệu PTBV",
        "scenario": "Yêu cầu cấp quyền truy cập tài liệu nội bộ / bảo mật",
        "risk": "Tài liệu bảo mật cao bị người ngoài ban xem trái phép hoặc ngược lại chuyên viên cần dữ liệu phục vụ kiểm toán bị chậm trễ cấp quyền.",
        "value": "Quản trị an toàn thông tin theo chuẩn phân quyền tài liệu nội bộ, có vết phê duyệt cấp quyền rõ ràng.",
        "trigger_recipient": "User bấm Xin quyền -> Gửi Người sở hữu tài liệu & Admin hệ thống",
        "template": "{user_name} ({department}) vừa gửi yêu cầu xin quyền xem/tải tài liệu nội bộ \"{doc_name}\" (Mã: {doc_code}). Vui lòng phê duyệt cấp quyền.",
        "type": "Yêu cầu xử lý (Action Required)",
        "priority": "Trung bình (P2)"
    }
]

for r_idx, sc in enumerate(proposed_details, 6):
    is_even = (r_idx % 2 == 0)
    row_fill = fill_zebra if is_even else fill_white
    ws3.row_dimensions[r_idx].height = 54
    
    ws3.cell(r_idx, 1, r_idx - 5).alignment = Alignment(horizontal="center", vertical="center")
    c2 = ws3.cell(r_idx, 2, sc["id"])
    c2.font = font_code_cell
    c2.alignment = Alignment(horizontal="center", vertical="center")
    
    c3 = ws3.cell(r_idx, 3, sc["module"])
    c3.font = font_bold_cell
    c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    c4 = ws3.cell(r_idx, 4, sc["scenario"])
    c4.font = font_bold_cell
    c4.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    ws3.cell(r_idx, 5, sc["risk"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws3.cell(r_idx, 6, sc["value"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws3.cell(r_idx, 7, sc["trigger_recipient"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws3.cell(r_idx, 8, sc["template"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    c9 = ws3.cell(r_idx, 9, sc["type"])
    c9.alignment = Alignment(horizontal="center", vertical="center")
    if "Action Required" in sc["type"]:
        c9.fill = FILL_REQ; c9.font = FONT_REQ
    elif "Warning" in sc["type"]:
        c9.fill = FILL_WARN; c9.font = FONT_WARN
    elif "Success" in sc["type"]:
        c9.fill = FILL_SUCC; c9.font = FONT_SUCC
    else:
        c9.fill = FILL_INFO; c9.font = FONT_INFO
        
    c10 = ws3.cell(r_idx, 10, sc["priority"])
    c10.alignment = Alignment(horizontal="center", vertical="center")
    if "P1" in sc["priority"]:
        c10.font = Font(name="Segoe UI", size=9, bold=True, color="991B1B")
    else:
        c10.font = Font(name="Segoe UI", size=9, bold=True, color="D97706")

    for col in range(1, 11):
        c = ws3.cell(r_idx, col)
        c.border = border_cell
        if col != 9:
            if c.font == font_data_cell or c.font is None:
                c.font = font_data_cell
            c.fill = row_fill

ws3.auto_filter.ref = f"A5:J{len(proposed_details) + 5}"
ws3.freeze_panes = "C6"

# =========================================================================
# SHEET 4: ĐẶC TẢ KỸ THUẬT UI & UX QUẢ CHUÔNG (NOTIFICATION SPECIFICATION)
# =========================================================================
ws4 = wb.create_sheet(title="4. Quy Chuẩn Kỹ Thuật UI-UX")
apply_title_block(
    ws4,
    title_text="ĐẶC TẢ KỸ THUẬT GIAO DIỆN & TRẢI NGHIỆM THÔNG BÁO (UI/UX SPECIFICATION)",
    subtitle_text="Quy chuẩn hiển thị biểu tượng Quả chuông, Popup danh sách thông báo, Phân loại trạng thái và Cơ chế tương tác Deep Link",
    col_span=6
)

headers4 = [
    "STT", "Hạng Mục Đặc Tả UI/UX", "Mô Tả Chi Tiết Yêu Cầu Giao Diện & Hành Vi",
    "Quy Tắc Kỹ Thuật (Technical Rules)", "Ví Dụ Minh Họa / Trạng Thái Hiển Thị", "Ghi Chú Triển Khai"
]

ws4.row_dimensions[5].height = 28
for col_idx, h_text in enumerate(headers4, 1):
    cell = ws4.cell(5, col_idx)
    cell.value = h_text
    cell.font = font_tbl_header
    cell.fill = fill_header
    cell.border = border_cell
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ui_specs = [
    {
        "item": "Biểu tượng Quả chuông (Header Bell Icon)",
        "desc": "Nằm ở góc trên bên phải thanh Header chính (cạnh avatar và tên người dùng). Hiển thị chấm đỏ hoặc badge số lượng khi có thông báo mới chưa đọc.",
        "rules": "- Khi có tin chưa đọc: Chấm đỏ góc trên (hoặc số 1-9, nếu >9 hiển thị 9+) có hiệu ứng pulse nhẹ (animate-pulse).\n- Khi không có tin chưa đọc: Biểu tượng quả chuông màu xám trung tính (text-black/45), hover chuyển màu VNA Deep Teal (#005F6E).",
        "example": "Icon Chuông + Badge đỏ có số \"3\" nhấp nháy nhẹ",
        "note": "Áp dụng tại component components/Layout.tsx"
    },
    {
        "item": "Popup Thông báo (Dropdown Panel)",
        "desc": "Mở rộng khi click vào Quả chuông. Kích thước tối ưu: rộng 380px - 420px, chiều cao tối đa 480px có thanh cuộn tinh tế (custom scrollbar).",
        "rules": "- Header popup: Chứa tiêu đề \"Thông báo\", số lượng chưa đọc và nút \"Đánh dấu tất cả là đã đọc\" (Check all as read).\n- Bộ lọc Tab: Gồm 2 tab: [Tất cả] và [Chưa đọc].\n- Danh sách tin: Xếp theo thời gian giảm dần (tin mới nhất ở trên cùng).",
        "example": "Panel thả xuống viền shadow-xl, nền trắng tinh khôi bo góc rounded-xl",
        "note": "Tự động đóng khi người dùng click ra ngoài (click-outside detector)"
    },
    {
        "item": "Thẻ Thông báo Đơn lẻ (Notification Item Card)",
        "desc": "Mỗi thông báo là một card tương tác: có icon phân loại mức độ, tiêu đề in đậm, trích đoạn nội dung, thời gian tương đối (relative time) và chấm xanh biểu thị trạng thái chưa đọc.",
        "rules": "- Trạng thái Chưa đọc (Unread): Nền xám nhạt/xanh băng (bg-blue-50/40), có chấm tròn xanh 6px ở góc phải, chữ in đậm hơn.\n- Trạng thái Đã đọc (Read): Nền trắng tinh, văn bản màu xám đen vừa phải (text-gray-700).\n- Hover effect: Đổi màu nền (hover:bg-gray-50) và hiển thị nút thao tác nhanh (xóa/đánh dấu đã đọc).",
        "example": "[Icon Yêu cầu] [Yêu cầu công bố số liệu] Chỉ tiêu E.1... (10 phút trước) [●]",
        "note": "Tối đa hiển thị 3 dòng nội dung, có text-truncate nếu quá dài"
    },
    {
        "item": "Mã Màu & Icon Phân Loại (Severity Mapping)",
        "desc": "Hệ thống icon và màu sắc trực quan giúp người dùng nhận diện ngay tính chất quan trọng của thông báo.",
        "rules": "1. Yêu cầu xử lý (Action Required): Icon AlertCircle đỏ (#DC2626), nền icon #FEE2E2.\n2. Cảnh báo (Warning): Icon AlertTriangle cam (#D97706), nền icon #FEF3C7.\n3. Thành công (Success): Icon CheckCircle2 xanh lá (#16A34A), nền icon #DCFCE7.\n4. Thông tin (Info): Icon Info xanh dương (#2563EB), nền icon #DBEAFE.",
        "example": "Đỏ = Cần duyệt ngay; Cam = Biểu đồ bị gỡ/thiếu data; Xanh lá = Đã xuất bản; Lam = Tài liệu mới",
        "note": "Dùng thư viện Lucide React icons đồng bộ hệ thống"
    },
    {
        "item": "Cơ Chế Deep Link & Điều Hướng (Action On Click)",
        "desc": "Khi click vào thông báo, hệ thống lập tức đánh dấu thông báo là 'Đã đọc' và tự động điều hướng người dùng thẳng tới màn hình mục tiêu kèm mở đúng tab và highlight dữ liệu.",
        "rules": "- Thông báo CMS Quản lý Biểu đồ: Điều hướng tới /cms-manage, tự động chuyển sang tab 'Quản lý biểu đồ công bố' và focus dòng biểu đồ tương ứng.\n- Thông báo Điều chỉnh số liệu: Điều hướng tới /publish-adjust, tự chọn chỉ tiêu và mở modal lịch sử phiên bản.\n- Thông báo Tài liệu: Điều hướng tới /documents, mở modal xem chi tiết file đính kèm.",
        "example": "Click thông báo Yêu cầu công bố -> Chuyển ngay đến CMS Tab Backlog và nhấp nháy thẻ chỉ tiêu",
        "note": "Truyền query params hoặc custom event qua window.dispatchEvent"
    },
    {
        "item": "Thông Báo Bổ Trợ: Toast & Push Alert",
        "desc": "Đối với các thông báo khẩn cấp hoặc thao tác trực tiếp của người dùng, hiển thị thêm Toast Popup nổi ở góc phải màn hình trong 4 giây.",
        "rules": "- Toast chỉ hiển thị tức thời cho người thực hiện thao tác hoặc thông báo khẩn cấp (P1).\n- Quả chuông lưu trữ bền vững (Persistent Notification) trong 30 ngày hoặc tối đa 100 thông báo gần nhất trên tài khoản.\n- Có thể tích hợp gửi Email digest hàng ngày/tuần đối với các yêu cầu phê duyệt chưa xử lý.",
        "example": "Toast xanh góc phải: 'Đã kích hoạt phiên bản v1.2 thành công!' (Tự ẩn sau 4s)",
        "note": "Tích hợp đồng bộ với cơ chế Toast hiện có của hệ thống VNA"
    },
    {
        "item": "Cơ Chế Cập Nhật Thời Gian Thực (Real-time Sync)",
        "desc": "Đảm bảo khi một tổ ban gửi yêu cầu hoặc Admin CMS thực hiện kéo thả biểu đồ, quả chuông của các bên liên quan cập nhật ngay mà không cần F5 tải lại trang.",
        "rules": "- Phía Client: Lắng nghe qua WebSocket hoặc Server-Sent Events (SSE). Dự phòng cơ chế Polling ngầm mỗi 30-60 giây.\n- Cập nhật số đếm badge tức thời và phát âm thanh nhẹ (tùy chọn trong Cài đặt thông báo).\n- Trong môi trường Prototype hiện tại: Sử dụng Custom Event 'vna_notification_event' và localStorage để sync giữa các tab trình duyệt.",
        "example": "Admin duyệt biểu đồ bên máy A -> Quả chuông máy B của Tổ ban tự nhảy số 1 đỏ",
        "note": "Đã sẵn sàng hook tích hợp trong Settings.tsx"
    }
]

for r_idx, sp in enumerate(ui_specs, 6):
    is_even = (r_idx % 2 == 0)
    row_fill = fill_zebra if is_even else fill_white
    ws4.row_dimensions[r_idx].height = 65
    
    ws4.cell(r_idx, 1, r_idx - 5).alignment = Alignment(horizontal="center", vertical="center")
    c2 = ws4.cell(r_idx, 2, sp["item"])
    c2.font = font_bold_cell
    c2.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    ws4.cell(r_idx, 3, sp["desc"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws4.cell(r_idx, 4, sp["rules"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws4.cell(r_idx, 5, sp["example"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws4.cell(r_idx, 6, sp["note"]).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for col in range(1, 7):
        c = ws4.cell(r_idx, col)
        c.border = border_cell
        if c.font == font_data_cell or c.font is None:
            c.font = font_data_cell
        c.fill = row_fill

ws4.auto_filter.ref = f"A5:F{len(ui_specs) + 5}"
ws4.freeze_panes = "C6"

# Apply column widths for all sheets
for ws in [ws1, ws2, ws3, ws4]:
    for col in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(5, ws.max_row + 1):
            val = ws.cell(row, col).value
            if val is not None:
                lines = str(val).split("\n")
                line_max = max(len(l) for l in lines) if lines else 0
                if line_max > max_len:
                    max_len = line_max
        header_val = ws.cell(5, col).value
        header_len = len(str(header_val)) if header_val else 10
        recommended = max(max_len * 1.05, header_len + 4)
        
        # Specific custom widths based on sheet
        if ws == ws4:
            if col == 1: ws.column_dimensions[col_letter].width = 7
            elif col == 2: ws.column_dimensions[col_letter].width = 24
            elif col == 3: ws.column_dimensions[col_letter].width = 38
            elif col == 4: ws.column_dimensions[col_letter].width = 44
            elif col == 5: ws.column_dimensions[col_letter].width = 32
            elif col == 6: ws.column_dimensions[col_letter].width = 28
        elif ws == ws3:
            if col == 1: ws.column_dimensions[col_letter].width = 7
            elif col == 2: ws.column_dimensions[col_letter].width = 16
            elif col == 3: ws.column_dimensions[col_letter].width = 24
            elif col == 4: ws.column_dimensions[col_letter].width = 30
            elif col == 5: ws.column_dimensions[col_letter].width = 40
            elif col == 6: ws.column_dimensions[col_letter].width = 38
            elif col == 7: ws.column_dimensions[col_letter].width = 34
            elif col == 8: ws.column_dimensions[col_letter].width = 46
            elif col == 9: ws.column_dimensions[col_letter].width = 20
            elif col == 10: ws.column_dimensions[col_letter].width = 16
        else:
            if col == 1: ws.column_dimensions[col_letter].width = 7
            elif col == 2: ws.column_dimensions[col_letter].width = 16
            elif col == 3: ws.column_dimensions[col_letter].width = 24
            elif col == 4: ws.column_dimensions[col_letter].width = 22
            elif col == 5: ws.column_dimensions[col_letter].width = 30
            elif col == 6: ws.column_dimensions[col_letter].width = 32
            elif col == 7: ws.column_dimensions[col_letter].width = 24
            elif col == 8: ws.column_dimensions[col_letter].width = 28
            elif col == 9: ws.column_dimensions[col_letter].width = 34
            elif col == 10: ws.column_dimensions[col_letter].width = 52
            elif col == 11: ws.column_dimensions[col_letter].width = 20
            elif col == 12: ws.column_dimensions[col_letter].width = 36
            elif col == 13: ws.column_dimensions[col_letter].width = 24
            elif col == 14: ws.column_dimensions[col_letter].width = 14

wb.save(excel_file)
print(f"Excel file successfully generated at: {excel_file}")
