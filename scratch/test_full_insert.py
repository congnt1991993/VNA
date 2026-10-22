# -*- coding: utf-8 -*-
import docx, copy, json

src_doc = 'images/document/BA DOC/URD/Backup/01. Tài liệu URD_v2.1_13.10.2026_backup_before_2_6.docx'
test_doc_path = 'scratch/test_urd_2_6.docx'

doc = docx.Document(src_doc)

# Find tbl_template
tbl_template = None
for child in doc.element.body:
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        p = docx.text.paragraph.Paragraph(child, doc)
        if '2.4.1.1' in p.text:
            tbl_found = True
            continue
    if 'tbl_found' in locals() and tbl_found and tag == 'tbl':
        tbl_template = docx.table.Table(child, doc)
        break

with open('scratch/section_2_6_data.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

def create_indicator_table(doc, ind):
    new_tbl_element = copy.deepcopy(tbl_template._tbl)
    tbl = docx.table.Table(new_tbl_element, doc)
    
    # Row 0: Code
    tbl.rows[0].cells[1].text = ''
    r0 = tbl.rows[0].cells[1].paragraphs[0].add_run(ind['code'])
    r0.bold = True
    
    # Row 1: Name
    tbl.rows[1].cells[1].text = ''
    r1 = tbl.rows[1].cells[1].paragraphs[0].add_run(f"{ind['name_en']} - {ind['name_vi']}")
    r1.italic = True
    
    # Row 2: Desc
    tbl.rows[2].cells[1].text = ''
    r2 = tbl.rows[2].cells[1].paragraphs[0].add_run(ind['desc'])
    r2.bold = False
    
    # Row 3: Pillar
    tbl.rows[3].cells[1].text = ''
    r3 = tbl.rows[3].cells[1].paragraphs[0].add_run(ind['pillar'])
    r3.bold = False
    
    # Row 4: Unit
    tbl.rows[4].cells[1].text = ''
    r4 = tbl.rows[4].cells[1].paragraphs[0].add_run(ind['unit'])
    r4.bold = False
    
    # Row 5: Header already "Mô tả chi tiết"
    # Row 6: Definition paragraphs
    tc6 = tbl.rows[6].cells[0]._tc
    p_elements = tc6.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p')
    for p_el in p_elements[1:]:
        tc6.remove(p_el)
    
    first_p = docx.text.paragraph.Paragraph(p_elements[0], doc)
    first_p.text = ''
    
    for idx, def_text in enumerate(ind['definition']):
        if idx == 0:
            target_p = first_p
        else:
            target_p = docx.text.paragraph.Paragraph(docx.oxml.OxmlElement('w:p'), doc)
            tc6.append(target_p._p)
        
        if def_text.startswith('1. ') or def_text.startswith('2. ') or def_text.startswith('3. '):
            colon_pos = def_text.find(':')
            if colon_pos != -1:
                prefix = def_text[:colon_pos + 1]
                body_text = def_text[colon_pos + 1:]
            else:
                prefix = def_text[:3]
                body_text = def_text[3:]
            r_pre = target_p.add_run(prefix)
            r_pre.bold = True
            r_body = target_p.add_run(body_text)
            r_body.bold = False
        else:
            target_p.add_run(def_text)
            
    return tbl

# Locate P[255] (2.6. Chỉ tiêu thuyêt) and P[259] (3. YÊU CẦU TÍCH HỢP)
p_2_6 = None
p_sec3 = None
for i, p in enumerate(doc.paragraphs):
    if '2.6. Chỉ tiêu thuyêt' in p.text:
        p_2_6 = p
    elif '3. YÊU CẦU TÍCH HỢP' in p.text and p.text.strip().startswith('3.'):
        p_sec3 = p
        break

print('Found p_2_6:', p_2_6.text if p_2_6 else None)
print('Found p_sec3:', p_sec3.text if p_sec3 else None)

# Update title of 2.6
p_2_6.text = "2.6. Chỉ tiêu thuyết minh"

# Remove the empty paragraphs between p_2_6 and p_sec3
body = doc.element.body
body_elements = list(body)
idx_2_6 = body_elements.index(p_2_6._p)
idx_sec3 = body_elements.index(p_sec3._p)
print(f'Elements between 2.6 and Section 3: {idx_sec3 - idx_2_6 - 1}')

for j in range(idx_2_6 + 1, idx_sec3):
    body.remove(body_elements[j])

# Helper to insert element before p_sec3
def insert_before_sec3(element):
    p_sec3._p.addprevious(element)

def add_p_before_sec3(text, style_name='Normal'):
    p = docx.text.paragraph.Paragraph(docx.oxml.OxmlElement('w:p'), doc)
    p.style = doc.styles[style_name]
    if text:
        p.add_run(text)
    insert_before_sec3(p._p)
    return p

# 1. Intro paragraph
intro_p = add_p_before_sec3("Mục này tổng hợp danh mục các chỉ tiêu thuyết minh (định tính, chính sách và cơ chế quản trị) thuộc bộ tiêu chuẩn GRI Standards 2021 và Sổ tay Báo cáo Phát triển bền vững Hàng không IATA ASRH Edition 2 (2023), được phân công quản lý và thu thập dữ liệu bởi Ban Kế hoạch phát triển và Ban Tổ chức nhân lực của Tổng công ty Hàng không Việt Nam.")
add_p_before_sec3("")

# 2. Section 2.6.1 Ban KHPT
add_p_before_sec3("2.6.1. Ban Kế hoạch phát triển (Trụ cột Quản trị)", "Heading 3")
add_p_before_sec3("")

for ind in data['khpt']:
    # Heading 4
    heading_text = f"{ind['num']}. Chỉ tiêu {ind['code']} - {ind['name_en']} - {ind['name_vi']}"
    add_p_before_sec3(heading_text, "Heading 4")
    # Table
    tbl = create_indicator_table(doc, ind)
    insert_before_sec3(tbl._tbl)
    # Spacing
    add_p_before_sec3("")

# 3. Section 2.6.2 Ban TCNL
add_p_before_sec3("2.6.2. Ban Tổ chức nhân lực", "Heading 3")
add_p_before_sec3("")

for ind in data['tcnl']:
    # Heading 4
    heading_text = f"{ind['num']}. Chỉ tiêu {ind['code']} - {ind['name_en']} - {ind['name_vi']}"
    add_p_before_sec3(heading_text, "Heading 4")
    # Table
    tbl = create_indicator_table(doc, ind)
    insert_before_sec3(tbl._tbl)
    # Spacing
    add_p_before_sec3("")

doc.save(test_doc_path)
print('Saved test doc successfully to', test_doc_path)
