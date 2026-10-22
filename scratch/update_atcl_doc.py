import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import shutil
import os

target_path = '/Users/congnguyen/Library/CloudStorage/OneDrive-Personal/Documents/Synodus/GOV/01.VNA/VNA/images/document/BA DOC/URD/Phụ lục Ban An Toàn Chất Lượng.docx'
backup_path = '/Users/congnguyen/Library/CloudStorage/OneDrive-Personal/Documents/Synodus/GOV/01.VNA/VNA/images/document/BA DOC/URD/Phụ lục Ban An Toàn Chất Lượng_backup.docx'

if not os.path.exists(backup_path):
    print(f"Creating backup: {backup_path}")
    shutil.copy2(target_path, backup_path)

doc = docx.Document(target_path)

def set_cell_font(cell, text, font_name="Times New Roman", font_size=10, bold=False, italic=False, color_rgb=(0x1F, 0x29, 0x37)):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)
    return p

print("Updating Table 6: Cảnh báo và workflow khi nhập liệu...")
tbl6 = doc.tables[6]

# Row 1: Nhập sai định dạng số liệu, chưa hoàn thành file Excel
set_cell_font(tbl6.cell(1, 0), "Nhập sai định dạng số liệu, chưa hoàn thành file Excel", font_size=10, bold=False)
set_cell_font(tbl6.cell(1, 1), 
    "Hệ thống tự động kích hoạt bộ tiền kiểm định dạng (Excel Validation Engine). Khi phát hiện lỗi (sai kiểu số, ngày tháng không đúng chuẩn DD/MM/YYYY, thiếu trường bắt buộc, hoặc chỉ số tiếng ồn ngoài danh mục 'Đạt'/'Không đạt'), hệ thống mở popup hiển thị bảng chi tiết lỗi: tọa độ ô (Cell Ref như B3, C5), tên cột, giá trị lỗi và thông điệp giải thích; đồng thời đánh dấu đỏ dòng lỗi tại màn hình xem trước và khóa nút \"Xác nhận Import\".", 
    font_size=9.5)
set_cell_font(tbl6.cell(1, 2), "Chuyên viên / Đầu mối nhập liệu Ban ATCL", font_size=9.5)
set_cell_font(tbl6.cell(1, 3), "Không áp dụng (Hệ thống tự động kiểm tra và chặn lỗi tại bước nhập liệu trước khi ghi vào CSDL)", font_size=9.5)
set_cell_font(tbl6.cell(1, 4), "Lỗi định dạng (Validation Error) / Chặn nạp dữ liệu (Import Blocked). Bản ghi giữ nguyên trạng thái Nháp, yêu cầu người dùng hiệu chỉnh lại file Excel theo đúng file mẫu.", font_size=9.5)

# Row 2: Popup nhập liệu/Màn hình thông báo
set_cell_font(tbl6.cell(2, 0), "Popup nhập liệu/Màn hình thông báo", font_size=10, bold=False)
set_cell_font(tbl6.cell(2, 1), 
    "Hiển thị popup modal xác nhận lưu dữ liệu (Toast/Modal thông báo thành công hoặc cảnh báo chưa lưu khi rời trang); Modal luồng phê duyệt (QuickApprovalActions / ApprovalModalComponent) cho phép gửi duyệt kèm ghi chú, phê duyệt hoặc từ chối bắt buộc nêu rõ lý do; hiển thị nhật ký phê duyệt (ApprovalLogTable) gồm thời gian, người thực hiện, hành động và ý kiến xử lý. Đồng thời hiển thị popup cảnh báo đối với các sự cố an toàn bay nghiêm trọng (mức độ Open/Investigating).", 
    font_size=9.5)
set_cell_font(tbl6.cell(2, 2), "Đầu mối nhập liệu Ban ATCL (thao tác Nhập liệu, Lưu nháp, Gửi duyệt) / Lãnh đạo Ban ATCL (thao tác Phê duyệt, Từ chối)", font_size=9.5)
set_cell_font(tbl6.cell(2, 3), "Lãnh đạo Ban ATCL (duyệt cấp đơn vị) / Ban KHPT (kiểm tra tổng hợp cấp Tổng công ty)", font_size=9.5)
set_cell_font(tbl6.cell(2, 4), "Chuyển trạng thái linh hoạt theo vòng đời dữ liệu: Nháp (Draft/Inactive) → Chờ phê duyệt (Pending) → Đã duyệt (Active/Approved) hoặc Từ chối (Rejected - hoàn trả kèm lý do để chỉnh sửa).", font_size=9.5)

print("Updating Table 12: Danh sách tồn tại (Defects & Issues Log)...")
tbl12 = doc.tables[12]

# Replace row 1
def fill_tbl12_row(row, code, content, level, solution, deadline, owner):
    set_cell_font(row.cells[0], code, font_size=9.5, bold=True, color_rgb=(0x00, 0x5F, 0x6E))
    set_cell_font(row.cells[1], content, font_size=9.5)
    level_color = (0xD9, 0x77, 0x06) if level == "Trung bình" else (0x4B, 0x55, 0x63)
    set_cell_font(row.cells[2], level, font_size=9.5, bold=True, color_rgb=level_color)
    set_cell_font(row.cells[3], solution, font_size=9.5)
    set_cell_font(row.cells[4], deadline, font_size=9.5)
    set_cell_font(row.cells[5], owner, font_size=9.5)

fill_tbl12_row(tbl12.rows[1],
    "DEF-ATCL-01",
    "Tích hợp tự động qua API với hệ thống quản lý an toàn VHAT (để tự động lấy số vụ chim va chạm, sét đánh, sự cố an toàn bay) chưa hoàn tất việc mở firewall và cấp tài khoản API trong giai đoạn UAT đợt 1.",
    "Trung bình",
    "Tạm thời sử dụng chức năng \"Nhập Excel\" định kỳ hàng tháng theo biểu mẫu chuẩn. Ban CĐSCN và đối tác kỹ thuật phối hợp với đội quản trị VHAT hoàn tất tích hợp API Gateway trong Giai đoạn 2.",
    "30/11/2026",
    "Ban CĐSCN, Ban ATCL, Đơn vị phát triển"
)

# Add row 2
row2 = tbl12.add_row()
fill_tbl12_row(row2,
    "DEF-ATCL-02",
    "Cần bổ sung cơ chế cảnh báo và cập nhật tự động khi Tổng công ty tiếp nhận hoặc thuê thêm tàu bay mới mà chưa có cấu hình chuẩn chứng chỉ tiếng ồn trong danh mục AMOS.",
    "Thấp",
    "Hệ thống tự động phát hiện số đăng ký tàu bay mới (Registration No.) chưa gán chuẩn tiếng ồn ICAO Annex 16 và hiển thị thông báo để đầu mối Ban ATCL cập nhật theo giấy chứng nhận AOC mới nhất.",
    "15/11/2026",
    "Ban ATCL, Đơn vị phát triển"
)

# Add row 3
row3 = tbl12.add_row()
fill_tbl12_row(row3,
    "DEF-ATCL-03",
    "Định dạng tệp PDF kết xuất báo cáo tiếng ồn (Airline E-1) và sự cố an toàn (GRI 403-2) cần tối ưu hóa độ phân giải biểu đồ Sankey và biểu đồ xu hướng theo khổ giấy A4 Landscape.",
    "Thấp",
    "Cập nhật cấu hình bộ sinh PDF (PDF Export Module) hỗ trợ định dạng linh hoạt theo mẫu báo cáo phát triển bền vững thường niên của VNA.",
    "31/10/2026",
    "Đơn vị phát triển"
)

print("Updating Table 10: Kết quả đối soát dữ liệu...")
tbl10 = doc.tables[10]
# Row 1: Airline E-1
set_cell_font(tbl10.cell(1, 0), "Airline E-1 (Tiếng ồn)", font_size=9.5, bold=True)
set_cell_font(tbl10.cell(1, 1), "Năm 2024 - 2026", font_size=9.5)
set_cell_font(tbl10.cell(1, 2), "Danh sách đội tàu bay và chứng nhận tiếng ồn từ AMOS đối chiếu với DWH", font_size=9.5)
set_cell_font(tbl10.cell(1, 3), "Khớp 100% (Toàn bộ đội tàu bay đạt chuẩn tiếng ồn ICAO Annex 16 Chapter 4/14)", font_size=9.5)
set_cell_font(tbl10.cell(1, 4), "0%", font_size=9.5)
set_cell_font(tbl10.cell(1, 5), "Đạt", font_size=9.5, bold=True, color_rgb=(0x16, 0xA3, 0x4A))

# Row 2: GRI 403-2
set_cell_font(tbl10.cell(2, 0), "GRI 403-2 (Sự cố an toàn)", font_size=9.5, bold=True)
set_cell_font(tbl10.cell(2, 1), "Năm 2024 - 2026", font_size=9.5)
set_cell_font(tbl10.cell(2, 2), "Số lượng sự cố thực tế từ VHAT đối chiếu với kết quả hiển thị trên hệ thống Netzero", font_size=9.5)
set_cell_font(tbl10.cell(2, 3), "Khớp 100% số liệu sự cố chim va chạm, sét đánh và sự cố khai thác theo từng tháng", font_size=9.5)
set_cell_font(tbl10.cell(2, 4), "0%", font_size=9.5)
set_cell_font(tbl10.cell(2, 5), "Đạt", font_size=9.5, bold=True, color_rgb=(0x16, 0xA3, 0x4A))

print("Updating Table 11: Kết quả phân quyền...")
tbl11 = doc.tables[11]
# Row 1: Đầu mối Ban ATCL
set_cell_font(tbl11.cell(1, 0), "Đầu mối Ban ATCL", font_size=9.5, bold=True)
set_cell_font(tbl11.cell(1, 1), "Nhập liệu chỉ số tiếng ồn (Airline E-1), phát thải KNK nội địa (Scope 1), nhập nhật ký sự cố an toàn (GRI 403-2), Import file Excel và gửi duyệt kỳ báo cáo", font_size=9.5)
set_cell_font(tbl11.cell(1, 2), "Đạt", font_size=9.5, bold=True, color_rgb=(0x16, 0xA3, 0x4A))
set_cell_font(tbl11.cell(1, 3), "Đã phân quyền đúng vai trò Đầu mối nhập liệu đơn vị Ban ATCL", font_size=9.5)

# Row 2: Ban Lãnh đạo
set_cell_font(tbl11.cell(2, 0), "Ban Lãnh đạo", font_size=9.5, bold=True)
set_cell_font(tbl11.cell(2, 1), "Xem dashboard báo cáo ATCL, rà soát số liệu đối chiếu, thực hiện Phê duyệt hoặc Từ chối kèm lý do", font_size=9.5)
set_cell_font(tbl11.cell(2, 2), "Đạt", font_size=9.5, bold=True, color_rgb=(0x16, 0xA3, 0x4A))
set_cell_font(tbl11.cell(2, 3), "Đã phân quyền đúng vai trò Lãnh đạo phê duyệt Ban ATCL", font_size=9.5)

print("Updating Table 7: Ma trận truy vết...")
tbl7 = doc.tables[7]
# Fill row 1 (YCQ-ATCL-01)
set_cell_font(tbl7.cell(1, 0), "YCQ-ATCL-01", font_size=9.5, bold=True, color_rgb=(0x00, 0x5F, 0x6E))
set_cell_font(tbl7.cell(1, 1), "Hình 01: Màn hình nhập liệu ATCL (Form & Import Excel)", font_size=9.5)
set_cell_font(tbl7.cell(1, 2), "Mức ồn (dB), Phát thải Scope 1 (tCO2e), Nhật ký sự cố", font_size=9.5)
set_cell_font(tbl7.cell(1, 3), "Kiểm thử giao diện nhập liệu, lưu nháp, kiểm tra lỗi định dạng và quy trình gửi duyệt.", font_size=9.5)

# Fill row 2 (YCQ-ATCL-02)
set_cell_font(tbl7.cell(2, 0), "YCQ-ATCL-02", font_size=9.5, bold=True, color_rgb=(0x00, 0x5F, 0x6E))
set_cell_font(tbl7.cell(2, 1), "Hình 01, Hình 04: Tích hợp dữ liệu nguồn (Excel, FIMS, VHAT)", font_size=9.5)
set_cell_font(tbl7.cell(2, 2), "Dữ liệu AMOS, FIMS, VHAT đồng bộ vào Data Warehouse", font_size=9.5)
set_cell_font(tbl7.cell(2, 3), "Kiểm thử nạp dữ liệu từ file Excel, kiểm tra log đồng bộ và cảnh báo dữ liệu thiếu/sai.", font_size=9.5)

# Add row 3 (YCQ-ATCL-03) if not present
if len(tbl7.rows) < 4:
    r3 = tbl7.add_row()
    set_cell_font(r3.cells[0], "YCQ-ATCL-03", font_size=9.5, bold=True, color_rgb=(0x00, 0x5F, 0x6E))
    set_cell_font(r3.cells[1], "Hình 02, Hình 03, Hình 04: Dashboard báo cáo ATCL", font_size=9.5)
    set_cell_font(r3.cells[2], "Biểu đồ tròn cơ cấu tiếng ồn, biểu đồ Sankey, biểu đồ cột sự cố an toàn", font_size=9.5)
    set_cell_font(r3.cells[3], "Kiểm thử hiển thị biểu đồ Metabase, tooltip hover, bộ lọc thời gian và xuất báo cáo PDF.", font_size=9.5)

doc.save(target_path)
print(f"Successfully updated and saved: {target_path}")
