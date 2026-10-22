import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

target_path = '/Users/congnguyen/Library/CloudStorage/OneDrive-Personal/Documents/Synodus/GOV/01.VNA/VNA/images/document/BA DOC/URD/Phụ lục Ban An Toàn Chất Lượng.docx'

doc = docx.Document(target_path)

def set_cell_font(cell, text, font_name="Times New Roman", font_size=10, bold=False, italic=False, color_rgb=(0x1F, 0x29, 0x37)):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)
    return p

# 1. Update Section 5 paragraph to explicitly state offline approval
print("Updating Section 5 paragraph (Offline approval note)...")
for i, p in enumerate(doc.paragraphs):
    if "5. Cảnh báo và workflow khi nhập liệu" in p.text:
        # Check next paragraph
        if i + 1 < len(doc.paragraphs):
            p_next = doc.paragraphs[i + 1]
            p_next.text = "Quy trình phê duyệt số liệu được thực hiện ngoại tuyến (offline) theo quy chế vận hành thực tế của Vietnam Airlines. Hệ thống phần mềm không áp dụng luồng phê duyệt trực tuyến (không có tính năng gửi duyệt, phê duyệt hay từ chối online); số liệu sau khi được người dùng kiểm tra định dạng và bấm Lưu (hoặc Import Excel thành công) sẽ được ghi nhận trực tiếp vào cơ sở dữ liệu Clean DB và cập nhật tức thời lên Dashboard báo cáo."
            p_next.runs[0].font.name = "Times New Roman"
            p_next.runs[0].font.size = Pt(10.5)
            p_next.runs[0].font.italic = True
            p_next.runs[0].font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        break

# 2. Update Table 6 (Cảnh báo khi nhập liệu - Không có luồng phê duyệt)
print("Updating Table 6 (Removing online approval workflow)...")
tbl6 = doc.tables[6]

# Row 1: Nhập sai định dạng số liệu, chưa hoàn thành file Excel
set_cell_font(tbl6.cell(1, 0), "Nhập sai định dạng số liệu, chưa hoàn thành file Excel", font_size=10, bold=False)
set_cell_font(tbl6.cell(1, 1), 
    "Hệ thống tự động kích hoạt bộ tiền kiểm định dạng (Excel Validation Engine). Khi phát hiện lỗi (sai kiểu số, ngày tháng không đúng chuẩn DD/MM/YYYY, thiếu trường bắt buộc, hoặc chỉ số tiếng ồn ngoài danh mục 'Đạt'/'Không đạt'), hệ thống mở popup hiển thị bảng chi tiết lỗi: tọa độ ô (Cell Ref như B3, C5), tên cột, giá trị lỗi và thông điệp giải thích; đồng thời đánh dấu đỏ dòng lỗi tại màn hình xem trước và khóa nút \"Xác nhận Import\".", 
    font_size=9.5)
set_cell_font(tbl6.cell(1, 2), "Chuyên viên / Đầu mối nhập liệu Ban ATCL", font_size=9.5)
set_cell_font(tbl6.cell(1, 3), "Không áp dụng (Hệ thống không có luồng phê duyệt online; việc kiểm soát và chặn lỗi được thực hiện tự động trước khi lưu vào CSDL)", font_size=9.5)
set_cell_font(tbl6.cell(1, 4), "Lỗi định dạng (Validation Error) / Chặn nạp dữ liệu (Import Blocked). Yêu cầu người dùng chỉnh sửa lại file Excel theo đúng file mẫu trước khi nạp vào hệ thống.", font_size=9.5)

# Row 2: Popup nhập liệu/Màn hình thông báo
set_cell_font(tbl6.cell(2, 0), "Popup nhập liệu/Màn hình thông báo", font_size=10, bold=False)
set_cell_font(tbl6.cell(2, 1), 
    "Hệ thống hiển thị popup modal thông báo kết quả lưu dữ liệu (Toast/Modal thông báo \"Đã lưu dữ liệu thành công!\" khi ghi nhận số liệu trực tiếp vào CSDL Clean DB; hoặc hiển thị cảnh báo xác nhận khi người dùng đóng/rời khỏi màn hình khi có dữ liệu chưa lưu). Đối với các chỉ số an toàn khai thác bay, hệ thống hiển thị badge trạng thái sự cố (Open, Investigating, Closed) và cảnh báo trực quan khi phát sinh sự cố nghiêm trọng hoặc khi các chỉ số tiếng ồn (dB), phát thải Scope 1 vượt ngưỡng định mức cho phép.", 
    font_size=9.5)
set_cell_font(tbl6.cell(2, 2), "Chuyên viên / Đầu mối nhập liệu Ban ATCL", font_size=9.5)
set_cell_font(tbl6.cell(2, 3), "Không áp dụng (Hệ thống không có luồng phê duyệt online; số liệu được ghi nhận trực tiếp vào CSDL Clean DB và cập nhật tức thời lên Dashboard sau khi lưu)", font_size=9.5)
set_cell_font(tbl6.cell(2, 4), "Đã lưu thành công (Saved / Active) — Dữ liệu được ghi nhận trực tiếp vào cơ sở dữ liệu Clean DB và cập nhật tức thời lên Dashboard báo cáo.", font_size=9.5)

# 3. Update Table 7 (Ma trận truy vết - Bỏ "gửi duyệt")
print("Updating Table 7 (Removing 'gửi duyệt')...")
tbl7 = doc.tables[7]
set_cell_font(tbl7.cell(1, 0), "YCQ-ATCL-01", font_size=9.5, bold=True, color_rgb=(0x00, 0x5F, 0x6E))
set_cell_font(tbl7.cell(1, 1), "Hình 01: Màn hình nhập liệu ATCL (Form & Import Excel)", font_size=9.5)
set_cell_font(tbl7.cell(1, 2), "Mức ồn (dB), Phát thải Scope 1 (tCO2e), Nhật ký sự cố", font_size=9.5)
set_cell_font(tbl7.cell(1, 3), "Kiểm thử giao diện nhập liệu, kiểm tra lỗi định dạng, lưu trữ trực tiếp vào CSDL Clean DB và cập nhật lên Dashboard.", font_size=9.5)

# 4. Update Table 11 (Phân quyền - Bỏ "gửi duyệt", "phê duyệt hoặc từ chối")
print("Updating Table 11 (Removing approval roles)...")
tbl11 = doc.tables[11]
# Row 1: Đầu mối Ban ATCL
set_cell_font(tbl11.cell(1, 0), "Đầu mối Ban ATCL", font_size=9.5, bold=True)
set_cell_font(tbl11.cell(1, 1), "Nhập liệu chỉ số tiếng ồn (Airline E-1), nhập liệu phát thải KNK nội địa (Scope 1), nhập nhật ký sự cố an toàn (GRI 403-2), Import file Excel và lưu trữ dữ liệu trực tiếp vào hệ thống", font_size=9.5)
set_cell_font(tbl11.cell(1, 2), "Đạt", font_size=9.5, bold=True, color_rgb=(0x16, 0xA3, 0x4A))
set_cell_font(tbl11.cell(1, 3), "Đã phân quyền đúng vai trò Đầu mối nhập liệu Ban ATCL (nhập liệu và lưu trực tiếp vào CSDL)", font_size=9.5)

# Row 2: Ban Lãnh đạo
set_cell_font(tbl11.cell(2, 0), "Ban Lãnh đạo", font_size=9.5, bold=True)
set_cell_font(tbl11.cell(2, 1), "Xem dashboard báo cáo ATCL, rà soát đối chiếu số liệu chỉ tiêu tiếng ồn, phát thải và sự cố an toàn trên hệ thống", font_size=9.5)
set_cell_font(tbl11.cell(2, 2), "Đạt", font_size=9.5, bold=True, color_rgb=(0x16, 0xA3, 0x4A))
set_cell_font(tbl11.cell(2, 3), "Đã phân quyền đúng vai trò Xem báo cáo và tra cứu dữ liệu của Ban Lãnh đạo (không có thao tác phê duyệt online)", font_size=9.5)

# 5. Update Section 4 paragraph to clarify role permissions without online approval
print("Updating Section 4 (Clarifying roles without online approval)...")
for i, p in enumerate(doc.paragraphs):
    if "4. Phân quyền" in p.text:
        # P49: - Lãnh đạo đơn vị
        if i + 1 < len(doc.paragraphs) and "Lãnh đạo đơn vị" in doc.paragraphs[i + 1].text:
            doc.paragraphs[i + 1].text = "- Lãnh đạo đơn vị: Quyền tra cứu, xem Dashboard báo cáo tổng hợp và theo dõi tiến độ các chỉ tiêu của Ban ATCL (phê duyệt số liệu thực hiện ngoại tuyến theo quy chế nội bộ của TCT)."
            doc.paragraphs[i + 1].runs[0].font.name = "Times New Roman"
            doc.paragraphs[i + 1].runs[0].font.size = Pt(10)
        # P50: - Đầu mối nhập liệu đơn vị
        if i + 2 < len(doc.paragraphs) and "Đầu mối nhập liệu đơn vị" in doc.paragraphs[i + 2].text:
            doc.paragraphs[i + 2].text = "- Đầu mối nhập liệu đơn vị: Quyền nhập liệu, chỉnh sửa, Import file Excel và lưu trữ trực tiếp số liệu vào cơ sở dữ liệu Clean DB."
            doc.paragraphs[i + 2].runs[0].font.name = "Times New Roman"
            doc.paragraphs[i + 2].runs[0].font.size = Pt(10)
        # P51: - Quản trị viên
        if i + 3 < len(doc.paragraphs) and "Quản trị viên" in doc.paragraphs[i + 3].text:
            doc.paragraphs[i + 3].text = "- Quản trị viên: Quản trị tài khoản người dùng, phân quyền truy cập và cấu hình tham số hệ thống."
            doc.paragraphs[i + 3].runs[0].font.name = "Times New Roman"
            doc.paragraphs[i + 3].runs[0].font.size = Pt(10)
        break

doc.save(target_path)
print(f"Successfully saved updated document without online approval workflow: {target_path}")
