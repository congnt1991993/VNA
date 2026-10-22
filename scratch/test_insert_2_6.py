# -*- coding: utf-8 -*-
import docx, copy, json
import lxml.etree

src_doc = 'images/document/BA DOC/URD/Backup/01. Tài liệu URD_v2.1_13.10.2026_backup_before_2_6.docx'
test_doc_path = 'scratch/test_urd_2_6.docx'

doc = docx.Document(src_doc)

# 1. Find template table from 2.4.1.1
tbl_template = None
for child in doc.element.body:
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        p = docx.text.paragraph.Paragraph(child, doc)
        if '2.4.1.1' in p.text:
            tbl_found = True
            continue
    if 'tbl_found' in locals() and tbl_found and tag == 'tbl':
        tbl_template = docx.table.Table(child, doc)
        break

print('Found tbl_template:', tbl_template is not None)

# Load data
with open('scratch/section_2_6_data.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

# Function to create an indicator table
def create_indicator_table(doc, ind):
    new_tbl_element = copy.deepcopy(tbl_template._tbl)
    tbl = docx.table.Table(new_tbl_element, doc)
    
    # Row 0: Code
    tbl.rows[0].cells[1].text = ''
    r0 = tbl.rows[0].cells[1].paragraphs[0].add_run(ind['code'])
    r0.bold = True
    
    # Row 1: Name
    tbl.rows[1].cells[1].text = ''
    r1 = tbl.rows[1].cells[1].paragraphs[0].add_run(f"{ind['name_en']} - {ind['name_vi']}")
    r1.italic = True
    
    # Row 2: Desc
    tbl.rows[2].cells[1].text = ''
    r2 = tbl.rows[2].cells[1].paragraphs[0].add_run(ind['desc'])
    r2.bold = False
    
    # Row 3: Pillar
    tbl.rows[3].cells[1].text = ''
    r3 = tbl.rows[3].cells[1].paragraphs[0].add_run(ind['pillar'])
    r3.bold = False
    
    # Row 4: Unit
    tbl.rows[4].cells[1].text = ''
    r4 = tbl.rows[4].cells[1].paragraphs[0].add_run(ind['unit'])
    r4.bold = False
    
    # Row 5: Header already "Mô tả chi tiết"
    # Row 6: Definition paragraphs
    tc6 = tbl.rows[6].cells[0]._tc
    p_elements = tc6.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p')
    for p_el in p_elements[1:]:
        tc6.remove(p_el)
    
    first_p = docx.text.paragraph.Paragraph(p_elements[0], doc)
    first_p.text = ''
    
    for idx, def_text in enumerate(ind['definition']):
        if idx == 0:
            target_p = first_p
        else:
            target_p = docx.text.paragraph.Paragraph(docx.oxml.OxmlElement('w:p'), doc)
            tc6.append(target_p._p)
        
        # Check if text starts with number or bold prefix
        if def_text.startswith('1. ') or def_text.startswith('2. ') or def_text.startswith('3. '):
            prefix = def_text[:def_text.find(':') + 1] if ':' in def_text else def_text[:3]
            body_text = def_text[len(prefix):]
            r_pre = target_p.add_run(prefix)
            r_pre.bold = True
            r_body = target_p.add_run(body_text)
            r_body.bold = False
        else:
            target_p.add_run(def_text)
            
    return tbl

# Test creating one table
t = create_indicator_table(doc, data['khpt'][0])
print('Table created successfully. Rows:', len(t.rows))
print('Row 0:', t.rows[0].cells[0].text, '|', t.rows[0].cells[1].text)
print('Row 6 paras:', len(t.rows[6].cells[0].paragraphs))
for p in t.rows[6].cells[0].paragraphs:
    print('  P:', p.text[:80])
